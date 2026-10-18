import subprocess, sys
from openpyxl import load_workbook


def _wb():
    subprocess.run([sys.executable, 'make_xlsx2.py'], check=True)
    return load_workbook('Planning_Stunfest2026_Indie.xlsx')


def test_friday_sheet_has_jeff_column():
    ws = _wb()['Vendredi 23 oct']
    assert 'Jeff' in [c.value for c in ws[2]]


def test_friday_briefing_cell():
    ws = _wb()['Vendredi 23 oct']
    header = [c.value for c in ws[2]]
    col = header.index('François') + 1
    assert ws.cell(row=3 + 2, column=1).value == '10h00'
    assert ws.cell(row=3 + 2, column=col).value == 'Briefing'


def _texts(ws):
    return [str(c.value) for row in ws.iter_rows() for c in row if c.value]


def test_notes_and_legend_updated():
    for ws in _wb().worksheets:
        texts = ' '.join(_texts(ws))
        assert '11h à 22h' in texts
        assert 'à partir de 20h' not in texts
    fri = ' '.join(_texts(_wb()['Vendredi 23 oct']))
    assert 'Briefing' in fri and '10h00-10h30' in fri


def test_no_vote_in_legend():
    for ws in _wb().worksheets:
        assert not any('Vote' in t for t in _texts(ws))
